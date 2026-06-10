from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import Pick, Result

PICKS_HEADERS = [
    "id",
    "created_at",
    "match_time",
    "league",
    "home",
    "away",
    "market",
    "line",
    "selection",
    "probability",
    "fair_odds",
    "book_odds",
    "edge",
    "est_goals",
    "status",
    "score_home",
    "score_away",
    "profit_units",
    "settled_at",
    "source",
    "notes",
]


class PickStore:
    def __init__(self, path: Path):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._ensure()

    def append_new(self, picks: list[Pick], now: datetime) -> list[Pick]:
        wb = load_workbook(self.path)
        ws = wb["Picks"]
        existing = {row[0].value for row in ws.iter_rows(min_row=2, max_col=1) if row[0].value}
        saved: list[Pick] = []
        for pick in picks:
            if pick.id in existing:
                continue
            ws.append(
                [
                    pick.id,
                    now.isoformat(timespec="seconds"),
                    pick.fixture.match_time.isoformat(timespec="seconds"),
                    pick.fixture.league,
                    pick.fixture.home,
                    pick.fixture.away,
                    pick.market,
                    pick.line,
                    pick.selection,
                    round(pick.probability, 4),
                    round(pick.fair_odds, 3),
                    pick.book_odds,
                    round(pick.edge, 4) if pick.edge is not None else None,
                    round(pick.est_goals, 3),
                    "PENDING",
                    None,
                    None,
                    None,
                    None,
                    pick.fixture.source,
                    pick.notes,
                ]
            )
            saved.append(pick)
        wb.save(self.path)
        return saved

    def settle(self, results: list[Result], now: datetime) -> int:
        wb = load_workbook(self.path)
        ws = wb["Picks"]
        headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
        settled = 0
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, headers["status"]).value != "PENDING":
                continue
            home = str(ws.cell(row, headers["home"]).value).upper()
            away = str(ws.cell(row, headers["away"]).value).upper()
            result = _find_result(results, home, away)
            if not result:
                continue
            total = result.total_goals
            line = float(ws.cell(row, headers["line"]).value)
            book_odds = ws.cell(row, headers["book_odds"]).value
            won = total > line
            profit = None
            if book_odds:
                profit = (float(book_odds) - 1) if won else -1
            ws.cell(row, headers["status"]).value = "WON" if won else "LOST"
            ws.cell(row, headers["score_home"]).value = result.home_goals
            ws.cell(row, headers["score_away"]).value = result.away_goals
            ws.cell(row, headers["profit_units"]).value = round(profit, 3) if profit is not None else None
            ws.cell(row, headers["settled_at"]).value = now.isoformat(timespec="seconds")
            settled += 1
        self._refresh_summary(wb)
        wb.save(self.path)
        return settled

    def summary(self) -> dict[str, float | int]:
        wb = load_workbook(self.path, data_only=True)
        ws = wb["Picks"]
        headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
        total = won = lost = pending = 0
        profit = 0.0
        for row in range(2, ws.max_row + 1):
            status = ws.cell(row, headers["status"]).value
            if not status:
                continue
            total += 1
            if status == "WON":
                won += 1
            elif status == "LOST":
                lost += 1
            elif status == "PENDING":
                pending += 1
            row_profit = ws.cell(row, headers["profit_units"]).value
            if isinstance(row_profit, (int, float)):
                profit += row_profit
        graded = won + lost
        return {
            "total": total,
            "won": won,
            "lost": lost,
            "pending": pending,
            "win_rate": (won / graded) if graded else 0,
            "profit_units": profit,
        }

    def _ensure(self) -> None:
        if self.path.exists():
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Picks"
        ws.append(PICKS_HEADERS)
        wb.create_sheet("Summary")
        wb.save(self.path)

    def _refresh_summary(self, wb) -> None:
        if "Summary" not in wb.sheetnames:
            wb.create_sheet("Summary")
        ws = wb["Summary"]
        ws.delete_rows(1, ws.max_row)
        summary = self.summary_from_workbook(wb)
        ws.append(["metric", "value"])
        for key, value in summary.items():
            ws.append([key, value])

    @staticmethod
    def summary_from_workbook(wb) -> dict[str, float | int]:
        ws = wb["Picks"]
        headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
        won = lost = pending = total = 0
        profit = 0.0
        for row in range(2, ws.max_row + 1):
            status = ws.cell(row, headers["status"]).value
            if not status:
                continue
            total += 1
            won += 1 if status == "WON" else 0
            lost += 1 if status == "LOST" else 0
            pending += 1 if status == "PENDING" else 0
            row_profit = ws.cell(row, headers["profit_units"]).value
            if isinstance(row_profit, (int, float)):
                profit += row_profit
        graded = won + lost
        return {
            "total_picks": total,
            "won": won,
            "lost": lost,
            "pending": pending,
            "win_rate": round(won / graded, 4) if graded else 0,
            "profit_units": round(profit, 3),
        }


def _find_result(results: list[Result], home: str, away: str) -> Result | None:
    for result in results:
        if result.home.upper() == home and result.away.upper() == away:
            return result
        if result.home.upper() == away and result.away.upper() == home:
            return Result(result.match_time, home, away, result.away_goals, result.home_goals, result.source)
    return None
