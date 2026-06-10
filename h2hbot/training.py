from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook

from .data_service import LEAGUES, MatchPick, _clean, _split_team_player
from .settings import Settings

TRAINING_HEADERS = [
    "id",
    "created_at",
    "league_id",
    "league",
    "match_time",
    "home_player",
    "away_player",
    "market",
    "pick",
    "probability",
    "value_score",
    "confidence",
    "sample_quality",
    "goal_trend_score",
    "winner_score_home",
    "winner_score_away",
    "draw_risk_score",
    "status",
    "score_home",
    "score_away",
    "settled_at",
    "risk",
    "argument",
]

SCORE_RE = re.compile(r"(?P<h>\d{1,2})\s*-\s*(?P<a>\d{1,2})")


class TrainingStore:
    def __init__(self, path: Path):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._ensure()

    def record_strong_picks(self, picks: list[MatchPick], now: datetime, min_value: float = 68.0) -> int:
        wb = load_workbook(self.path)
        ws = wb["Training"]
        headers = _headers(ws)
        existing = {row[0].value for row in ws.iter_rows(min_row=2, max_col=1) if row[0].value}
        added = 0
        for pick in picks:
            if not _is_trackable_pick(pick, min_value):
                continue
            if pick.id in existing:
                continue
            ws.append(
                [
                    pick.id,
                    now.isoformat(timespec="seconds"),
                    pick.league_id,
                    pick.league,
                    pick.iso_time,
                    pick.home_player,
                    pick.away_player,
                    pick.best_market,
                    pick.best_pick,
                    pick.best_probability,
                    pick.value_score,
                    pick.confidence,
                    pick.sample_quality,
                    pick.goal_trend_score,
                    pick.winner_score_home,
                    pick.winner_score_away,
                    pick.draw_risk_score,
                    "PENDING",
                    None,
                    None,
                    None,
                    pick.risk,
                    pick.argument,
                ]
            )
            existing.add(pick.id)
            added += 1
        self._refresh_summary(wb, headers)
        wb.save(self.path)
        return added

    def settle(self, results: dict[str, dict[str, tuple[int, int]]], now: datetime) -> int:
        wb = load_workbook(self.path)
        ws = wb["Training"]
        headers = _headers(ws)
        settled = 0
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, headers["status"]).value != "PENDING":
                continue
            league_id = str(ws.cell(row, headers["league_id"]).value)
            home = _clean(str(ws.cell(row, headers["home_player"]).value))
            away = _clean(str(ws.cell(row, headers["away_player"]).value))
            score = results.get(league_id, {}).get(_match_key(home, away))
            if not score:
                continue
            home_score, away_score = score
            pick = str(ws.cell(row, headers["pick"]).value)
            status = _grade_pick(pick, home_score, away_score)
            ws.cell(row, headers["status"]).value = status
            ws.cell(row, headers["score_home"]).value = home_score
            ws.cell(row, headers["score_away"]).value = away_score
            ws.cell(row, headers["settled_at"]).value = now.isoformat(timespec="seconds")
            settled += 1
        self._refresh_summary(wb, headers)
        wb.save(self.path)
        return settled

    def summary(self) -> dict:
        wb = load_workbook(self.path, data_only=True)
        ws = wb["Training"]
        headers = _headers(ws)
        rows = _training_rows(ws, headers)
        return _summary_from_rows(rows)

    def _ensure(self) -> None:
        if not self.path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "Training"
            ws.append(TRAINING_HEADERS)
            wb.create_sheet("Summary")
            wb.save(self.path)
            return
        wb = load_workbook(self.path)
        if "Training" not in wb.sheetnames:
            ws = wb.create_sheet("Training", 0)
            ws.append(TRAINING_HEADERS)
        if "Summary" not in wb.sheetnames:
            wb.create_sheet("Summary")
        ws = wb["Training"]
        current = [cell.value for cell in ws[1]]
        missing = [header for header in TRAINING_HEADERS if header not in current]
        for header in missing:
            ws.cell(1, ws.max_column + 1).value = header
        wb.save(self.path)

    def _refresh_summary(self, wb, headers: dict[str, int]) -> None:
        ws = wb["Training"]
        rows = _training_rows(ws, headers)
        summary = _summary_from_rows(rows)
        out = wb["Summary"]
        out.delete_rows(1, out.max_row)
        out.append(["metric", "value"])
        for key, value in summary["overall"].items():
            out.append([key, value])
        out.append([])
        out.append(["league_id", "market", "total", "won", "lost", "pending", "win_rate"])
        for row in summary["groups"]:
            out.append(row)


def fetch_results() -> dict[str, dict[str, tuple[int, int]]]:
    results: dict[str, dict[str, tuple[int, int]]] = {}
    for league in LEAGUES:
        try:
            html = requests.get(league["totalcorner_url"], headers={"User-Agent": "Mozilla/5.0"}, timeout=20).text
            results[league["id"]] = _parse_results_table(html)
        except Exception:  # noqa: BLE001
            results[league["id"]] = {}
    return results


def _parse_results_table(html: str) -> dict[str, tuple[int, int]]:
    soup = BeautifulSoup(html, "html.parser")
    tables = soup.select("table")
    if not tables:
        return {}
    results: dict[str, tuple[int, int]] = {}
    for row in tables[-1].select("tr"):
        cols = [c.get_text(" ", strip=True) for c in row.find_all("td")]
        if len(cols) < 5 or cols[1].strip().lower() != "full":
            continue
        home_team, home_player = _split_team_player(cols[2])
        away_team, away_player = _split_team_player(cols[4])
        score = SCORE_RE.search(cols[3])
        if not home_player or not away_player or not score:
            continue
        results[_match_key(home_player, away_player)] = (int(score.group("h")), int(score.group("a")))
    return results


def _is_trackable_pick(pick: MatchPick, min_value: float) -> bool:
    if pick.value_score < min_value:
        return False
    if pick.best_pick.startswith("No Bet"):
        return False
    if pick.sample_quality != "Alta":
        return False
    if "moderado" in pick.best_pick.lower():
        return False
    if pick.league_id == "h2hgg":
        return pick.best_market == "Over/Under 2.5"
    return pick.best_market == "1X2"


def _grade_pick(pick: str, home_score: int, away_score: int) -> str:
    total = home_score + away_score
    if pick.startswith("Over 2.5"):
        return "WON" if total > 2.5 else "LOST"
    if pick == "Under 2.5":
        return "WON" if total < 2.5 else "LOST"
    if pick == "Local gana":
        return "WON" if home_score > away_score else "LOST"
    if pick == "Visita gana":
        return "WON" if away_score > home_score else "LOST"
    return "VOID"


def _headers(ws) -> dict[str, int]:
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1]) if cell.value}


def _training_rows(ws, headers: dict[str, int]) -> list[dict]:
    rows = []
    for row in range(2, ws.max_row + 1):
        item = {name: ws.cell(row, col).value for name, col in headers.items()}
        if item.get("id"):
            rows.append(item)
    return rows


def _summary_from_rows(rows: list[dict]) -> dict:
    total = len(rows)
    won = sum(1 for row in rows if row.get("status") == "WON")
    lost = sum(1 for row in rows if row.get("status") == "LOST")
    pending = sum(1 for row in rows if row.get("status") == "PENDING")
    graded = won + lost
    groups: dict[tuple[str, str], dict[str, int]] = {}
    for row in rows:
        key = (str(row.get("league_id")), str(row.get("market")))
        groups.setdefault(key, {"total": 0, "won": 0, "lost": 0, "pending": 0})
        groups[key]["total"] += 1
        status = row.get("status")
        if status == "WON":
            groups[key]["won"] += 1
        elif status == "LOST":
            groups[key]["lost"] += 1
        elif status == "PENDING":
            groups[key]["pending"] += 1
    group_rows = []
    for (league_id, market), values in sorted(groups.items()):
        group_graded = values["won"] + values["lost"]
        win_rate = round(values["won"] / group_graded, 4) if group_graded else 0
        group_rows.append([league_id, market, values["total"], values["won"], values["lost"], values["pending"], win_rate])
    return {
        "overall": {
            "total": total,
            "won": won,
            "lost": lost,
            "pending": pending,
            "win_rate": round(won / graded, 4) if graded else 0,
        },
        "groups": group_rows,
    }


def _match_key(home: str, away: str) -> str:
    return f"{_clean(home)}|{_clean(away)}"


def run_training_cycle(settings: Settings, dashboard: dict | None = None) -> dict:
    from .data_service import load_dashboard

    now = datetime.now(ZoneInfo(settings.timezone))
    if dashboard is None:
        dashboard = load_dashboard(settings)
    picks = [MatchPick(**item) for item in dashboard["matches"]]
    store = TrainingStore(settings.picks_xlsx)
    added = store.record_strong_picks(picks, now)
    settled = store.settle(fetch_results(), now)
    return {"added": added, "settled": settled, "summary": store.summary()}
